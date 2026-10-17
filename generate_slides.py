import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Color palette
    BG_DARK = RGBColor(11, 25, 44)         # Deep Navy #0B192C
    BG_LIGHT = RGBColor(248, 250, 252)     # Slate 50 #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)      # White
    PRIMARY_BLUE = RGBColor(0, 92, 238)    # GCash Blue #005CEE
    ACCENT_TEAL = RGBColor(13, 148, 136)   # Teal #0D9488
    TEXT_DARK = RGBColor(15, 23, 42)       # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139)   # Slate 500
    TEXT_WHITE = RGBColor(255, 255, 255)
    
    GREEN_SAFE = RGBColor(16, 185, 129)    # Emerald 500
    AMBER_CAUTION = RGBColor(245, 158, 11) # Amber 500
    RED_SUSPICIOUS = RGBColor(239, 68, 68) # Red 500
    
    BORDER_LIGHT = RGBColor(226, 232, 240)
    
    blank_layout = prs.slide_layouts[6]
    
    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="IT 021 - USABLE SECURITY | GROUP 7"):
        # Category label
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_BLUE
        
        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 1: Title & The Hook
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, BG_DARK)
    
    # Top Tag
    tb1_tag = slide1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(0.4))
    tf1_tag = tb1_tag.text_frame
    p1_tag = tf1_tag.paragraphs[0]
    p1_tag.text = "TECHNOLOGICAL INSTITUTE OF THE PHILIPPINES • COLLEGE OF COMPUTER STUDIES • IT 021"
    p1_tag.font.size = Pt(11)
    p1_tag.font.bold = True
    p1_tag.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
    
    # THE HOOK (Large Banner Box)
    hook_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.3), Inches(11.3), Inches(1.6))
    hook_box.fill.solid()
    hook_box.fill.fore_color.rgb = RGBColor(23, 42, 69)
    hook_box.line.color.rgb = RGBColor(56, 189, 248)
    hook_box.line.width = Pt(1.5)
    
    tf_hook = hook_box.text_frame
    tf_hook.word_wrap = True
    tf_hook.margin_left = Inches(0.4)
    tf_hook.margin_right = Inches(0.4)
    tf_hook.margin_top = Inches(0.25)
    
    p_hlabel = tf_hook.paragraphs[0]
    p_hlabel.text = "THE OPENING REALITY:"
    p_hlabel.font.size = Pt(10)
    p_hlabel.font.bold = True
    p_hlabel.font.color.rgb = RGBColor(251, 191, 36) # Amber
    
    p_htext = tf_hook.add_paragraph()
    p_htext.text = "\"How a ₱5 printed sticker can drain your e-wallet in under 3 seconds.\""
    p_htext.font.size = Pt(22)
    p_htext.font.bold = True
    p_htext.font.color.rgb = TEXT_WHITE
    
    # Project Title & Subtitle
    tb_mtitle = slide1.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(8.0), Inches(1.8))
    tf_mtitle = tb_mtitle.text_frame
    tf_mtitle.word_wrap = True
    p_mt1 = tf_mtitle.paragraphs[0]
    p_mt1.text = "TapTignan"
    p_mt1.font.size = Pt(44)
    p_mt1.font.bold = True
    p_mt1.font.color.rgb = TEXT_WHITE
    
    p_mt2 = tf_mtitle.add_paragraph()
    p_mt2.text = "A Pre-Payment Verification and Reporting Interface for Filipino Mobile Wallet Users"
    p_mt2.font.size = Pt(16)
    p_mt2.font.color.rgb = RGBColor(148, 163, 184)
    
    # Team Info Box
    team_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.1), Inches(11.3), Inches(1.7))
    team_box.fill.solid()
    team_box.fill.fore_color.rgb = RGBColor(15, 30, 50)
    team_box.line.color.rgb = RGBColor(30, 58, 95)
    
    tf_team = team_box.text_frame
    tf_team.word_wrap = True
    tf_team.margin_left = Inches(0.4)
    tf_team.margin_top = Inches(0.2)
    
    p_t1 = tf_team.paragraphs[0]
    p_t1.text = "SUBMITTED BY: Group 7 (BSIT - IT31S4)  |  INSTRUCTOR: Allan A. Burgos"
    p_t1.font.size = Pt(11)
    p_t1.font.bold = True
    p_t1.font.color.rgb = RGBColor(56, 189, 248)
    
    p_t2 = tf_team.add_paragraph()
    p_t2.text = "• Mike Angelo Salamat (Project Lead & Prototype Developer)      • Adrian Carlo Quiambao (Technical & Prototype Lead)\n• Francheska Ann Nicole Rabanzo (Security & Docs Lead)          • Marwan Bherooze Sabino (Docs & Security Design)\n• Joana Marie Sostea (User & Usability Analyst)"
    p_t2.font.size = Pt(11)
    p_t2.font.color.rgb = RGBColor(226, 232, 240)
    
    # Add Shield Logo if exists
    shield_path = "assets/logos/taptignan-inspect-shield.png"
    if os.path.exists(shield_path):
        slide1.shapes.add_picture(shield_path, Inches(9.8), Inches(3.2), width=Inches(2.3))

    # -------------------------------------------------------------
    # SLIDE 2: The Problem: "Quishing" & Zero-Friction Dilemma
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "The Real-World Threat: The QR 'Quishing' Epidemic in the Philippines")
    
    # 3 Threat Cards
    col_w = Inches(3.64)
    col_gap = Inches(0.39)
    y_cards = Inches(1.6)
    card_h = Inches(4.2)
    
    threats = [
        ("T-01: Quishing Stickers", "Physical Scam Overlays", 
         "Attackers print counterfeit QR stickers and paste them over legitimate merchant counter stands at food stalls, MRT stations, and parking kiosks.",
         "Impact: The victim thinks they paid the cashier, but their money went straight to an attacker's disposable wallet.",
         RED_SUSPICIOUS),
        ("T-03: Merchant Spoofing", "Deceptive Display Names",
         "The physical poster says 'Cafe Manila', but the QR code resolves to an unrelated individual GCash account ('Juan Dela Cruz').",
         "Impact: Mobile wallets display only nominal text without verifying against registered corporate business directories.",
         AMBER_CAUTION),
        ("The Zero-Friction Dilemma", "Instant & Irreversible",
         "Mobile wallets prioritize speed. Zero friction means zero inspection. Transactions complete in 3 seconds with no safety gate.",
         "Impact: Once a user confirms, funds cannot be recalled, leaving victims with zero recourse and total balance loss.",
         PRIMARY_BLUE)
    ]
    
    for i, (t_code, t_title, t_desc, t_impact, t_color) in enumerate(threats):
        x = Inches(0.8) + i * (col_w + col_gap)
        c = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_cards, col_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_LIGHT
        c.line.width = Pt(1)
        
        # Color bar top
        bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_cards, col_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = t_color
        bar.line.fill.background()
        
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)
        
        p1 = tf.paragraphs[0]
        p1.text = t_code.upper()
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = t_color
        
        p2 = tf.add_paragraph()
        p2.text = t_title
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        
        p3 = tf.add_paragraph()
        p3.text = "\n" + t_desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = RGBColor(71, 85, 105)
        
        p4 = tf.add_paragraph()
        p4.text = "\n" + t_impact
        p4.font.size = Pt(11)
        p4.font.bold = True
        p4.font.color.rgb = TEXT_DARK
        
    # Bottom callout
    bot = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9))
    bot.fill.solid()
    bot.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bot.line.color.rgb = RGBColor(199, 210, 254)
    tf_b = bot.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0.3)
    p_bot = tf_b.paragraphs[0]
    p_bot.text = "CORE RESEARCH PRINCIPLE: User Problem → Security Risk → Usability Issue → Design Requirement → Prototype → Verification"
    p_bot.font.size = Pt(10)
    p_bot.font.bold = True
    p_bot.font.color.rgb = PRIMARY_BLUE

    # -------------------------------------------------------------
    # SLIDE 3: BEFORE VS AFTER TABLE (GCash vs. GCash + TapTignan)
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "Comparative Evaluation: Standard GCash vs. GCash + TapTignan")
    
    # Table dimensions
    t_left = Inches(0.8)
    t_top = Inches(1.5)
    t_width = Inches(11.7)
    t_height = Inches(5.4)
    
    rows = 7
    cols = 3
    table_shape = slide3.shapes.add_table(rows, cols, t_left, t_top, t_width, t_height)
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.5)  # Feature
    tbl.columns[1].width = Inches(4.4)  # Standard GCash
    tbl.columns[2].width = Inches(4.8)  # GCash + TapTignan
    
    headers = ["SECURITY INTERACTION", "STANDARD GCASH (BASELINE)", "GCASH + TAPTIGNAN (REDESIGN)"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(15, 23, 42)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.alignment = PP_ALIGN.LEFT
        
    table_data = [
        ("QR Scanning Guidance", 
         "Plain camera viewfinder; no alignment guidance, framing cues, or contextual security alerts.",
         "Active TapTignan alignment frame, dynamic hint text, camera permission guidance, and quick scenario testing."),
        ("Pre-Payment Verification",
         "Bypassed completely. Proceeds straight from scan to payment confirmation with zero security gate.",
         "Mandatory Pre-Payment Gate (FR-05). Automated 6-rule engine evaluates QR before funds can move."),
        ("Merchant Identity Check",
         "Displays nominal unverified text; vulnerable if scammer labels their personal account as 'Cafe Manila'.",
         "Validates accountId against verified business registry; detects and flags impersonation and mismatches instantly."),
        ("Micro-Merchants / Sari-Sari",
         "No distinction between corporate and personal wallets; leaves buyers confused or unguided.",
         "3-Tier Calibrated Verdict (D-3). Recognizes micro-merchants without warning fatigue; injects verbal confirmation prompt."),
        ("Suspicious Transaction Handling",
         "Zero risk flags. If a code is malicious, money is transferred irreversibly without friction.",
         "High-Risk Red Verdict (FR-07, FR-10). Plain-language finding explanations, 1-tap Cancel (safe exit), and 2-step override."),
        ("Community Reporting & Blacklist",
         "Buried multi-step support ticket form; slow, tedious, and often abandoned by users (DF-02).",
         "Frictionless 4-Click Reporting (FR-13). Predefined categories, instant Reference ID, and real-time local blacklisting.")
    ]
    
    for r_idx, (feat, gcash, taptignan) in enumerate(table_data, start=1):
        bg_col = CARD_BG if r_idx % 2 == 1 else RGBColor(241, 245, 249)
        
        c0 = tbl.cell(r_idx, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = bg_col
        c0.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = c0.text_frame.paragraphs[0]
        p0.text = feat
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_DARK
        
        c1 = tbl.cell(r_idx, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = bg_col
        c1.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = c1.text_frame.paragraphs[0]
        p1.text = "❌ " + gcash
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = RGBColor(100, 116, 139)
        
        c2 = tbl.cell(r_idx, 2)
        c2.fill.solid()
        c2.fill.fore_color.rgb = bg_col
        c2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p2 = c2.text_frame.paragraphs[0]
        p2.text = "✅ " + taptignan
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(15, 23, 42)

    # -------------------------------------------------------------
    # SLIDE 4: Architecture & The 3-Tier Calibrated Verdict
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "Core Innovation: The 3-Tier Calibrated Verdict (Decision D-3)")
    
    # 3 Tier Cards
    w4 = Inches(3.64)
    gap4 = Inches(0.39)
    y4 = Inches(1.5)
    h4 = Inches(5.3)
    
    tiers = [
        ("SAFE VERDICT", "Accredited Corporate Merchant",
         "GREEN SHIELD (WCAG AA)",
         "• QR format parses correctly (R-01)\n• Zero external embedded links (R-02)\n• Account matches corporate registry (R-03)\n• Confirmed accredited business (R-04)\n• No prior community scam reports (R-05)",
         "Action: 'Continue to Payment' (Confidence)\nExample: Cafe Manila Inc. (₱850.00)",
         GREEN_SAFE, RGBColor(236, 253, 245)),
        ("CAUTION VERDICT", "Micro-Merchant / Sari-Sari Store",
         "AMBER SHIELD (WCAG AA)",
         "• Solves Philippine Warning Fatigue!\n• 80%+ daily transactions are personal P2P\n• Valid format & unflagged, but personal wallet\n• Does NOT block or terrify the buyer\n• Injects Actionable Verbal Prompt:",
         "Verbal Check: \"Kayo po ba si Ate Rose?\"\nExample: Rosemarie Natividad (₱45.00)",
         AMBER_CAUTION, RGBColor(254, 243, 199)),
        ("SUSPICIOUS VERDICT", "Critical Security Risk Detected",
         "RED SHIELD (WCAG AA)",
         "• Immediate trigger on high-severity risk\n• Embedded phishing link (R-02)\n• Merchant name mismatch / spoofing (R-03)\n• Community reported scam sticker (R-05)\n• Primary Action: 'Cancel Payment'",
         "Safe Exit: 1-Tap Cancel (₱0 deducted)\nOverride: Requires 2-step explicit modal",
         RED_SUSPICIOUS, RGBColor(254, 226, 226))
    ]
    
    for i, (tag, subtitle, badge, bullets, footer, border_col, bg_col) in enumerate(tiers):
        x = Inches(0.8) + i * (w4 + gap4)
        c = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y4, w4, h4)
        c.fill.solid()
        c.fill.fore_color.rgb = bg_col
        c.line.color.rgb = border_col
        c.line.width = Pt(1.5)
        
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)
        
        p1 = tf.paragraphs[0]
        p1.text = tag
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = border_col
        
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        
        p3 = tf.add_paragraph()
        p3.text = badge
        p3.font.size = Pt(9.5)
        p3.font.bold = True
        p3.font.color.rgb = RGBColor(100, 116, 139)
        
        p4 = tf.add_paragraph()
        p4.text = "\n" + bullets
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = RGBColor(51, 65, 85)
        
        p5 = tf.add_paragraph()
        p5.text = "\n" + footer
        p5.font.size = Pt(10.5)
        p5.font.bold = True
        p5.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 5: Frictionless Reporting & Duress Mode
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "Community Intelligence & Anti-Coercion Security")
    
    # Left Box: Reporting Flow
    w5 = Inches(5.65)
    h5 = Inches(5.3)
    y5 = Inches(1.5)
    
    c_left = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y5, w5, h5)
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = CARD_BG
    c_left.line.color.rgb = BORDER_LIGHT
    
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.35)
    
    pl1 = tf_l.paragraphs[0]
    pl1.text = "FRICTIONLESS REPORTING ENGINE (FR-13 to FR-15)"
    pl1.font.size = Pt(11)
    pl1.font.bold = True
    pl1.font.color.rgb = PRIMARY_BLUE
    
    pl2 = tf_l.add_paragraph()
    pl2.text = "Solving DF-02 (Report Fatigue) & DF-03 (Black Hole)"
    pl2.font.size = Pt(15)
    pl2.font.bold = True
    pl2.font.color.rgb = TEXT_DARK
    
    pl3 = tf_l.add_paragraph()
    pl3.text = "\n1. 4 Structured Predefined Reasons:\n   • Recipient name does not match merchant\n   • QR contains suspicious link / phishing URL\n   • Forced or pressured into paying\n   • Unrecognized recipient or business\n\n2. Instant Transparent Tracking:\n   • Generates distinct Reference ID (TT-2026-XXXXX)\n   • Activity tab tracks status: 'Under Review' → 'Resolved'\n\n3. Real-Time Community Defense:\n   • Submitting a report immediately blacklists the raw QR and accountId in the local threat store, shielding all future scans!"
    pl3.font.size = Pt(11)
    pl3.font.color.rgb = RGBColor(51, 65, 85)
    
    # Right Box: Anti-Coercion Duress Mode
    c_right = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), y5, w5, h5)
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = RGBColor(254, 242, 242) # Soft Red
    c_right.line.color.rgb = RGBColor(254, 202, 202)
    
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = Inches(0.35)
    
    pr1 = tf_r.paragraphs[0]
    pr1.text = "SPECIAL DEFENSE: ANTI-COERCION PROTECTION"
    pr1.font.size = Pt(11)
    pr1.font.bold = True
    pr1.font.color.rgb = RED_SUSPICIOUS
    
    pr2 = tf_r.add_paragraph()
    pr2.text = "Duress MPIN & Decoy Balance (Threat T-05)"
    pr2.font.size = Pt(15)
    pr2.font.bold = True
    pr2.font.color.rgb = TEXT_DARK
    
    pr3 = tf_r.add_paragraph()
    pr3.text = "\n• The Coercion Reality:\n  In holdups or extortion, victims are forced to unlock their phones under direct threat of physical harm.\n\n• The Usable-Security Countermeasure:\n  - Normal Registered MPIN ('1234'):\n    Displays real wallet balance: ₱100,367.00\n  - Duress MPIN ('9999'):\n    Silently activates Decoy Mode displaying only ₱48.50!\n\n• Life-Safety Principle:\n  The app appears completely normal to satisfy attackers while safeguarding the user's life savings and flagging coercion."
    pr3.font.size = Pt(11)
    pr3.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # SLIDE 6: Prototype Implementation & Empirical Verification
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_header(slide6, "Technical Architecture & Empirical Verification")
    
    # 4 Metric Cards
    w6 = Inches(2.7)
    gap6 = Inches(0.3)
    y6_top = Inches(1.5)
    h6_top = Inches(2.2)
    
    metrics = [
        ("TEST SUITE", "29 / 29 PASSED", "100% Pass Rate across Auth, Security, Rules, and Reports", GREEN_SAFE),
        ("EXECUTION TIME", "2.2 SECONDS", "Satisfies NFR-01 (< 3s total verification time)", PRIMARY_BLUE),
        ("ARCHITECTURE", "REACT NATIVE", "Expo SDK 57, TypeScript, Zustand, Expo Router", ACCENT_TEAL),
        ("STANDARDS", "WCAG AA", "Icon + Color + Plain English & Tagalog labels", AMBER_CAUTION)
    ]
    
    for i, (m_tag, m_val, m_desc, m_col) in enumerate(metrics):
        x = Inches(0.8) + i * (w6 + gap6)
        c = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y6_top, w6, h6_top)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_LIGHT
        
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        
        p1 = tf.paragraphs[0]
        p1.text = m_tag
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = m_col
        
        p2 = tf.add_paragraph()
        p2.text = m_val
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        
        p3 = tf.add_paragraph()
        p3.text = m_desc
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = TEXT_MUTED

    # Bottom Implementation Table
    y6_bot = Inches(3.9)
    h6_bot = Inches(2.9)
    bot_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y6_bot, Inches(11.7), h6_bot)
    bot_box.fill.solid()
    bot_box.fill.fore_color.rgb = CARD_BG
    bot_box.line.color.rgb = BORDER_LIGHT
    
    tf_bb = bot_box.text_frame
    tf_bb.word_wrap = True
    tf_bb.margin_left = tf_bb.margin_right = tf_bb.margin_top = Inches(0.3)
    
    pbb1 = tf_bb.paragraphs[0]
    pbb1.text = "DESIGN FINDINGS RESOLUTION (TRACEABILITY MATRIX)"
    pbb1.font.size = Pt(11)
    pbb1.font.bold = True
    pbb1.font.color.rgb = PRIMARY_BLUE
    
    pbb2 = tf_bb.add_paragraph()
    pbb2.text = "• DF-01 (Users overlook suspicious recipient names): Resolved via mandatory pre-payment gate (FR-05) + explicit side-by-side merchant comparison cards (FR-08).\n• DF-02 (Reporting abandoned due to complex forms): Resolved via 4 single-tap categorized radio buttons + auto-populated QR payload details (FR-13).\n• DF-03 (Users don't know what happens after reporting): Resolved via unique tracking ID (TT-2026-XXXXX) and live Under Review / Resolved status in Activity tab (FR-14, FR-15).\n• Offline Capability & Privacy (NFR-04, NFR-05): 100% self-contained prototype with zero mock PII, local storage persistence, and zero remote data leakage."
    pbb2.font.size = Pt(10.5)
    pbb2.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # SLIDE 7: Conclusion & Live Demonstration
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_DARK)
    
    # Header on Dark
    tb7_h = slide7.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
    tf7_h = tb7_h.text_frame
    p7_tag = tf7_h.paragraphs[0]
    p7_tag.text = "USABLE SECURITY DEFENSE • CONCLUSION & DEMO"
    p7_tag.font.size = Pt(11)
    p7_tag.font.bold = True
    p7_tag.font.color.rgb = RGBColor(56, 189, 248)
    
    p7_title = tf7_h.add_paragraph()
    p7_title.text = "Empowering Filipino E-Wallet Users with Usable Defense"
    p7_title.font.size = Pt(24)
    p7_title.font.bold = True
    p7_title.font.color.rgb = TEXT_WHITE
    
    # Left Card: Live Demo Steps
    w7 = Inches(5.65)
    h7 = Inches(4.5)
    y7 = Inches(1.6)
    
    c7_l = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y7, w7, h7)
    c7_l.fill.solid()
    c7_l.fill.fore_color.rgb = RGBColor(15, 30, 50)
    c7_l.line.color.rgb = RGBColor(30, 58, 95)
    
    tf7_l = c7_l.text_frame
    tf7_l.word_wrap = True
    tf7_l.margin_left = tf7_l.margin_right = tf7_l.margin_top = Inches(0.3)
    
    pdl1 = tf7_l.paragraphs[0]
    pdl1.text = "LIVE DEMONSTRATION SCENARIOS"
    pdl1.font.size = Pt(11)
    pdl1.font.bold = True
    pdl1.font.color.rgb = RGBColor(251, 191, 36)
    
    pdl2 = tf7_l.add_paragraph()
    pdl2.text = "Running on Actual Tested Prototype:\n\n1. Scenario 1: Safe Merchant (Cafe Manila)\n   • Verifies accredited corporate registry → Green Shield\n\n2. Scenario 2: Micro-Merchant (Ate Rose Sari-Sari)\n   • Triggers Caution Tier with verbal prompt: \"Kayo po ba si Ate Rose?\"\n\n3. Scenario 3: Quishing Scam Sticker (MRT-3 Shaw)\n   • Blocks counterfeit QR → 1-Tap Cancel (₱0 lost) → 1-Tap Report\n\n4. Scenario 4: Duress MPIN ('9999')\n   • Unlocks Decoy Balance (₱48.50) under coercion"
    pdl2.font.size = Pt(10.5)
    pdl2.font.color.rgb = RGBColor(226, 232, 240)
    
    # Right Card: Closing Takeaways
    c7_r = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), y7, w7, h7)
    c7_r.fill.solid()
    c7_r.fill.fore_color.rgb = RGBColor(15, 30, 50)
    c7_r.line.color.rgb = RGBColor(30, 58, 95)
    
    tf7_r = c7_r.text_frame
    tf7_r.word_wrap = True
    tf7_r.margin_left = tf7_r.margin_right = tf7_r.margin_top = Inches(0.3)
    
    pdr1 = tf7_r.paragraphs[0]
    pdr1.text = "CORE DEFENSE TAKEAWAY"
    pdr1.font.size = Pt(11)
    pdr1.font.bold = True
    pdr1.font.color.rgb = RGBColor(56, 189, 248)
    
    pdr2 = tf7_r.add_paragraph()
    pdr2.text = "\n\"Security does not have to destroy convenience, and convenience must not compromise human safety.\"\n\nBy embedding calibrated pre-payment checks, contextual verbal prompts, and frictionless community reporting, TapTignan transforms passive victims into protected, vigilant users.\n\n---------------------------------------------\nMaraming salamat po!\nWe are now open for your questions."
    pdr2.font.size = Pt(12)
    pdr2.font.color.rgb = RGBColor(226, 232, 240)
    
    # Footer
    tb7_foot = slide7.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4))
    tf7_foot = tb7_foot.text_frame
    p7_f = tf7_foot.paragraphs[0]
    p7_f.text = "TapTignan Project • Group 7 • BSIT - IT31S4 • Technological Institute of the Philippines"
    p7_f.font.size = Pt(10)
    p7_f.font.color.rgb = RGBColor(100, 116, 139)
    p7_f.alignment = PP_ALIGN.CENTER
    
    output_path = "TapTignan_Defense_Group7.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
